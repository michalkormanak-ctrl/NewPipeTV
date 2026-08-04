from pathlib import Path

from docx import Document

from app.export.docx_exporter import to_docx
from tests.factories import sample_legislative_package as _sample_package


def test_docx_contains_title_headings_and_tables(tmp_path: Path) -> None:
    package = _sample_package()
    output_path = to_docx(package, tmp_path / "balik.docx")
    assert output_path.exists()

    document = Document(str(output_path))

    heading_texts = [
        p.text for p in document.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"
    ]
    assert any("A. Manažérske zhrnutie" in h for h in heading_texts)
    assert any("J. Kontrolná správa" in h for h in heading_texts)
    assert any("K. Otvorené otázky" in h for h in heading_texts)
    assert any("L. Zdroje" in h for h in heading_texts)

    # Tabuľka nálezov (J) - 1 hlavička + 2 riadky nálezov (rozdelené cez 2 severity skupiny)
    tables = document.tables
    assert len(tables) >= 1
    findings_table = tables[0]
    assert findings_table.rows[0].cells[0].text == "Kontrola"

    # Zdroje (L) - tabuľka s citáciou
    sources_table = tables[-1]
    assert "500/2022 Z. z." in sources_table.rows[1].cells[0].text


def test_docx_omits_empty_optional_section_heading(tmp_path: Path) -> None:
    package = _sample_package()
    package.paragraph_wording = None
    output_path = to_docx(package, tmp_path / "balik2.docx")
    document = Document(str(output_path))
    heading_texts = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert not any("E. Paragrafové znenie" in h for h in heading_texts)


def test_docx_preserves_paragraph_breaks_as_separate_paragraphs(tmp_path: Path) -> None:
    package = _sample_package()
    package.consolidated_text = "§ 12\n(1) Pôvodný text.\n(2) Nový text."
    output_path = to_docx(package, tmp_path / "balik3.docx")
    document = Document(str(output_path))
    body_texts = [p.text for p in document.paragraphs]
    assert "(1) Pôvodný text." in body_texts
    assert "(2) Nový text." in body_texts
