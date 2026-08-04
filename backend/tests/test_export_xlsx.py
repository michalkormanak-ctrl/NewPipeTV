from pathlib import Path

from openpyxl import load_workbook

from app.export.package import CommentExportRow
from app.export.xlsx_exporter import to_comments_xlsx, to_findings_xlsx
from tests.factories import sample_legislative_package


def test_findings_xlsx_has_header_and_rows(tmp_path: Path) -> None:
    package = sample_legislative_package()
    output_path = to_findings_xlsx(package, tmp_path / "kontrola.xlsx")
    assert output_path.exists()

    workbook = load_workbook(output_path)
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["Závažnosť", "Kontrola", "Miesto", "Zistenie", "Odporúčanie", "Stav"]
    assert sheet.max_row == 1 + len(package.control_findings)
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert any("§ 99" in str(row) for row in rows)


def test_comments_xlsx_maps_all_fields(tmp_path: Path) -> None:
    comments = [
        CommentExportRow(
            author="Ministerstvo financií SR",
            target_provision_label="§ 5 ods. 2",
            comment_type="zasadna",
            is_fundamental=True,
            text="Navrhujeme predĺžiť lehotu.",
            proposed_wording="30 dní sa nahrádza slovami 60 dní.",
            justification="Nedostatočný čas na vybavenie.",
            thematic_category="procesné lehoty",
            legal_argument="Zásada primeranosti.",
            result="ciastocne_akceptovana",
            resolution_method="Skrátenie na 45 dní po rozporovom konaní.",
            final_wording="45 dní",
        )
    ]
    output_path = to_comments_xlsx(comments, tmp_path / "pripomienky.xlsx")
    workbook = load_workbook(output_path)
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert "Autor" in header
    assert "Zásadná" in header
    row = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert row[0] == "Ministerstvo financií SR"
    assert row[3] == "áno"
    assert row[9] == "ciastocne_akceptovana"
