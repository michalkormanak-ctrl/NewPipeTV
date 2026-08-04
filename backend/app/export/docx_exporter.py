"""DOCX export (sekcia 15 zadania - musí zachovať nadpisy, číslovanie,
paragrafy, odseky, poznámky a tabuľky)."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.export.package import SECTION_TITLES, SEVERITY_LABELS, SEVERITY_ORDER, LegislativePackage


def _add_text_block(document: DocumentType, text: str) -> None:
    """Zachová štruktúru odsekov/viet ako samostatné Word odseky (nie jeden
    zlepený blok textu) - dôležité pre čitateľnosť paragrafového znenia."""
    for line in text.split("\n"):
        line = line.strip()
        if line:
            document.add_paragraph(line)


def build_docx(package: LegislativePackage) -> DocumentType:
    document = Document()

    title_paragraph = document.add_heading(package.title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    disclaimer = document.add_paragraph()
    disclaimer_run = disclaimer.add_run(package.disclaimer)
    disclaimer_run.bold = True
    disclaimer_run.italic = True

    meta = document.add_paragraph()
    as_of_note = " (predvolený - dátum nebol zadaný)" if package.as_of_was_default else ""
    meta.add_run(
        f"Vytvorené: {package.generated_at.isoformat()} · "
        f"Právny stav posudzovaný k: {package.as_of_date.isoformat()}{as_of_note}"
    ).italic = True

    for attr, heading in SECTION_TITLES:
        value = getattr(package, attr)
        if not value:
            continue
        document.add_heading(heading, level=1)
        _add_text_block(document, value)

    document.add_heading("J. Kontrolná správa", level=1)
    grouped = package.findings_by_severity()
    any_findings = False
    for severity in SEVERITY_ORDER:
        findings = grouped.get(severity, [])
        if not findings:
            continue
        any_findings = True
        document.add_heading(SEVERITY_LABELS[severity], level=2)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        header[0].text, header[1].text, header[2].text = "Kontrola", "Miesto", "Zistenie"
        for finding in findings:
            row = table.add_row().cells
            row[0].text = finding.check
            row[1].text = finding.location or ""
            row[2].text = finding.message
    if not any_findings:
        document.add_paragraph("Žiadne nálezy.")

    document.add_heading("K. Otvorené otázky", level=1)
    if package.open_questions:
        for question in package.open_questions:
            document.add_paragraph(question, style="List Number")
    else:
        document.add_paragraph("Žiadne otvorené otázky vyžadujúce rozhodnutie človeka.")

    document.add_heading("L. Zdroje", level=1)
    if package.sources:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        header[0].text, header[1].text, header[2].text = "Predpis", "Ustanovenie", "Zdroj"
        for source in package.sources:
            row = table.add_row().cells
            row[0].text = source.instrument_full_citation
            row[1].text = source.provision_label or ""
            row[2].text = source.source_url or ""
    else:
        document.add_paragraph("Žiadne zdroje neboli priložené.")

    return document


def to_docx(package: LegislativePackage, path: str | Path) -> Path:
    document = build_docx(package)
    output_path = Path(path)
    document.save(output_path)
    return output_path


def to_docx_bytes(package: LegislativePackage) -> bytes:
    buffer = io.BytesIO()
    build_docx(package).save(buffer)
    return buffer.getvalue()
